"""Generate PVC catalog migration SQL from legacy Excel workbook."""
from __future__ import annotations

from pathlib import Path

import openpyxl

EXCEL_PATH = Path(r"c:\Users\pmech\Desktop\Quotes\PVC Strip Curtains\AJ Baker\PVC Strip Door 30-4-24.xlsx")
OUT_PATH = Path(__file__).resolve().parents[1] / "lib" / "db" / "migrations" / "074_pvc_calculator_excel_catalog_sync.sql"
BRACKETS_PATH = Path(__file__).resolve().parents[1] / "lib" / "db" / "migrations" / "076_pvc_brackets_hanger_width_column.sql"


def esc(value: object) -> str:
    if value is None:
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def product_id_expr(product_code: str) -> str:
    return f"(SELECT id FROM public.pvc_products WHERE product_code = {esc(product_code)} LIMIT 1)"


HANGER_BRACKET_SKU: dict[int, str] = {
    100: "PVHZ010120",
    150: "PVHZ010125",
    200: "PVHZ010130",
    300: "PVHZ010135",
    400: "PVHZ010135",
}


def main() -> None:
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["Product 1"]

    products: dict[str, dict] = {}
    strip_specs: list[dict] = []
    headrails: dict[str, str] = {}
    brackets: list[dict] = []
    fittings: list[dict] = []

    grade_map = {
        "s": "standard",
        "p": "polar",
        "sf": "standard_freezer",
        "pf": "polar_freezer",
        "c": "colour",
    }
    profile_map = {"p": "smooth", "r": "ribbed", "c": "colour", "s": "sheet"}
    headrail_map = {"ss": "stainless", "gl": "galvanized", "al": "aluminium", "as": "plastic"}

    for row in range(112, 213):
        code = ws[f"M{row}"].value
        if not code:
            continue
        code = str(code).strip()
        width = ws[f"A{row}"].value
        thick = ws[f"B{row}"].value
        grade = ws[f"C{row}"].value
        profile = ws[f"D{row}"].value
        cost = ws[f"Q{row}"].value
        desc = ws[f"N{row}"].value
        if code not in products:
            products[code] = {
                "code": code,
                "name": desc,
                "cost": cost,
                "sell": cost,
            }
        if width and thick and grade and profile:
            spec_key = (
                int(width),
                float(thick),
                grade_map.get(str(grade), str(grade)),
                profile_map.get(str(profile), str(profile)),
            )
            strip_specs.append(
                {
                    "spec_key": spec_key,
                    "product_code": code,
                    "width": spec_key[0],
                    "thick": spec_key[1],
                    "grade": spec_key[2],
                    "surface": spec_key[3],
                }
            )

    for row in range(214, 224):
        code = ws[f"M{row}"].value
        if not code:
            continue
        code = str(code).strip()
        headrail_key = ws[f"A{row}"].value
        cost = ws[f"Q{row}"].value
        desc = ws[f"N{row}"].value
        if code not in products:
            products[code] = {"code": code, "name": desc, "cost": cost, "sell": cost}
        if headrail_key in headrail_map:
            headrails[headrail_map[headrail_key]] = code

    for row in range(225, 258):
        code = ws[f"M{row}"].value
        if not code:
            continue
        code = str(code).strip()
        cost = ws[f"Q{row}"].value
        desc = ws[f"N{row}"].value
        hanger_width = ws[f"A{row}"].value
        headrail_key = ws[f"E{row}"].value
        if cost and cost > 0:
            if code not in products:
                products[code] = {"code": code, "name": desc, "cost": cost, "sell": cost}
            if hanger_width:
                brackets.append(
                    {
                        "headrail": headrail_map.get(str(headrail_key), str(headrail_key)),
                        "hanger_width": int(hanger_width),
                        "product_code": code,
                    }
                )

    for row in range(259, 345):
        code = ws[f"M{row}"].value
        if not code:
            continue
        code = str(code).strip()
        cost = ws[f"Q{row}"].value
        desc = ws[f"N{row}"].value
        hanger_width = ws[f"A{row}"].value
        headrail_key = ws[f"E{row}"].value
        rivets = ws[f"I{row}"].value
        if cost and cost > 0:
            if code not in products:
                products[code] = {"code": code, "name": desc, "cost": cost, "sell": cost}
            if hanger_width and rivets:
                fittings.append(
                    {
                        "headrail": headrail_map.get(str(headrail_key), str(headrail_key)),
                        "hanger_width": int(hanger_width),
                        "product_code": code,
                        "rivets_per_strip": int(rivets),
                    }
                )

    products.setdefault("PACZ010150", {"code": "PACZ010150", "name": "Cardboard Tube 1800mm", "cost": 3.81, "sell": 3.81})

    lines: list[str] = [
        "-- Migration: sync PVC calculator catalog from Excel PVC Strip Door 30-4-24",
        "-- Generated by scripts/generate_pvc_excel_migration.py",
        "-- Resolves product_id by product_code so existing pvc_products rows keep their UUIDs.",
        "",
    ]

    for product in products.values():
        cost = product["cost"]
        if cost is None or (isinstance(cost, (int, float)) and cost <= 0):
            continue
        lines.append(
            f"UPDATE public.pvc_products SET cost_price = {cost}, sell_price = {product['sell']} "
            f"WHERE product_code = {esc(product['code'])};"
        )
    lines.append("")

    # DB unique_strip_spec is on (width, thickness, grade, surface) — one product per combo.
    specs_by_unique: dict[tuple, dict] = {}
    for spec in strip_specs:
        specs_by_unique[spec["spec_key"]] = spec

    for spec in specs_by_unique.values():
        pid = product_id_expr(spec["product_code"])
        lines.extend(
            [
                "INSERT INTO public.pvc_strip_specs (product_id, strip_width_mm, thickness_mm, material_grade, surface_type)",
                f"SELECT {pid}, {spec['width']}, {spec['thick']}, {esc(spec['grade'])}, {esc(spec['surface'])}",
                f"WHERE {pid} IS NOT NULL",
                "ON CONFLICT (strip_width_mm, thickness_mm, material_grade, surface_type)",
                "DO UPDATE SET product_id = EXCLUDED.product_id;",
                "",
            ]
        )

    brackets_by_width: dict[int, str] = dict(HANGER_BRACKET_SKU)
    for bracket in brackets:
        brackets_by_width[bracket["hanger_width"]] = bracket["product_code"]

    fittings_by_width: dict[int, dict] = {}
    for fitting in fittings:
        width = fitting["hanger_width"]
        existing = fittings_by_width.get(width)
        if existing is None or fitting["headrail"] in ("galvanized", "stainless"):
            fittings_by_width[width] = fitting

    lines.append("-- Headrails: optional if pvc_headrails.headrail_type exists.")
    lines.append("")

    for headrail_type, product_code in headrails.items():
        pid = product_id_expr(product_code)
        lines.extend(
            [
                "INSERT INTO public.pvc_headrails (headrail_type, product_id)",
                f"SELECT {esc(headrail_type)}, {pid}",
                f"WHERE {pid} IS NOT NULL",
                "  AND NOT EXISTS (",
                "    SELECT 1 FROM public.pvc_headrails h",
                f"    WHERE h.headrail_type = {esc(headrail_type)}",
                "  );",
                "",
                "UPDATE public.pvc_headrails h",
                f"SET product_id = {pid}",
                f"WHERE h.headrail_type = {esc(headrail_type)}",
                f"  AND {pid} IS NOT NULL;",
                "",
            ]
        )

    bracket_lines: list[str] = [
        "-- Migration: hanger_width_mm on pvc_brackets + Excel bracket product mapping",
        "-- Generated by scripts/generate_pvc_excel_migration.py",
        "-- Run after 074 (prices + strip specs).",
        "",
        "ALTER TABLE IF EXISTS public.pvc_brackets",
        "  ADD COLUMN IF NOT EXISTS hanger_width_mm INTEGER;",
        "",
        "CREATE UNIQUE INDEX IF NOT EXISTS idx_pvc_brackets_hanger_width_mm",
        "  ON public.pvc_brackets (hanger_width_mm)",
        "  WHERE hanger_width_mm IS NOT NULL;",
        "",
    ]

    for hanger_width, product_code in sorted(brackets_by_width.items()):
        pid = product_id_expr(product_code)
        bracket_lines.extend(
            [
                "INSERT INTO public.pvc_brackets (hanger_width_mm, product_id)",
                f"SELECT {hanger_width}, {pid}",
                f"WHERE {pid} IS NOT NULL",
                "  AND NOT EXISTS (",
                "    SELECT 1 FROM public.pvc_brackets b",
                f"    WHERE b.hanger_width_mm = {hanger_width}",
                "  );",
                "",
                "UPDATE public.pvc_brackets b",
                f"SET product_id = {pid}",
                f"WHERE b.hanger_width_mm = {hanger_width}",
                f"  AND {pid} IS NOT NULL;",
                "",
            ]
        )

    for hanger_width, rivets in sorted(
        {w: f["rivets_per_strip"] for w, f in fittings_by_width.items()}.items()
    ):
        lines.append(
            f"INSERT INTO public.pvc_calculator_settings (setting_key, setting_value) "
            f"VALUES ('rivets_per_hanger_{hanger_width}', {rivets}) "
            f"ON CONFLICT (setting_key) DO UPDATE SET setting_value = EXCLUDED.setting_value;"
        )
    lines.append("")

    pack_pid = product_id_expr("PACZ010150")
    lines.extend(
        [
            "INSERT INTO public.pvc_packaging (product_id)",
            f"SELECT {pack_pid}",
            f"WHERE {pack_pid} IS NOT NULL",
            "  AND NOT EXISTS (SELECT 1 FROM public.pvc_packaging);",
            "",
        ]
    )

    OUT_PATH.write_text("\n".join(lines), encoding="utf-8")
    BRACKETS_PATH.write_text("\n".join(bracket_lines), encoding="utf-8")
    print(f"Wrote {OUT_PATH}")
    print(f"Wrote {BRACKETS_PATH} ({len(brackets_by_width)} bracket widths)")


if __name__ == "__main__":
    main()
